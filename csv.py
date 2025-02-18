import os
import re
import docx
import xlwt

def extract_text_from_cv(docx_path):
    """
    Extracts text from a CV document.
    """
    doc = docx.Document(docx_path)           
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def extract_emails_and_contacts(text):
    """
    Extracts email IDs and contact numbers from text using regular expressions.
    """
    emails = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
    contacts = re.findall(r'\b(?:\d[ -]?)?\(?\d{3}\)?[- ]?\d{3}[- ]?\d{4}\b', text)
    return emails, contacts

def create_excel_file(data, output_path):
    """
    Creates an Excel file (.xls) from the extracted data.
    """
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet('CV Data')

    # Write headers
    sheet.write(0, 0, 'Email')
    sheet.write(0, 1, 'Contact')
    sheet.write(0, 2, 'Overall Text')

    # Write data
    for idx, row in enumerate(data, start=1):
        sheet.write(idx, 0, row['email'])
        sheet.write(idx, 1, row['contact'])
        sheet.write(idx, 2, row['text'])

    # Save the Excel file
    workbook.save(output_path)

def process_cvs(input_folder, output_path):
    """
    Processes all CVs in the input folder and creates an Excel file with extracted data.
    """
    cv_data = []

    for filename in os.listdir(input_folder):
        if filename.endswith(".docx"):
            cv_path = os.path.join(input_folder, filename)
            text = extract_text_from_cv(cv_path)
            emails, contacts = extract_emails_and_contacts(text)
            cv_data.append({
                'email': ", ".join(emails),
                'contact': ", ".join(contacts),
                'text': text
            })

    create_excel_file(cv_data, output_path)

# Main function
if __name__ == "__main__":
    input_folder = "path/to/input_folder"  # Update with the path to the folder containing CVs
    output_path = "output.xls"  # Path to save the output Excel file
    process_cvs(input_folder, output_path)
